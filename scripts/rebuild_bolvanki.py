#!/usr/bin/env python3
"""
Пересборка болванок data/templates/*.docx на базе шаблона инженера.

Старые болванки задают другие поля страницы (top≈1843), из‑за этого при склейке
таблицы отчётов (tblInd для top≈284) уезжают к левому краю.

Сохраняем только текст заголовка «Отчёт … за» из каждой болванки.
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from sk_reporter.paths import templates_dir

DEFAULT_ENGINEER_TEMPLATE = ROOT / "data" / "engineer" / "report_template.docx"


def _title_lines_from_bolvanka(path: Path) -> list[str]:
    doc = Document(str(path))
    lines = [para.text for para in doc.paragraphs]
    if not lines:
        return [""]
    if len(lines) == 1:
        return lines
    # Как в старых болванках: строка заголовка + пустая строка.
    if lines[1].strip() == "":
        return [lines[0], ""]
    return lines


def _clear_body_keep_sectpr(doc: Document) -> None:
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is sect_pr:
            continue
        body.remove(child)


def rebuild_bolvanka(
    engineer_template: Path,
    old_bolvanka: Path,
    output_path: Path,
) -> list[str]:
    title_lines = _title_lines_from_bolvanka(old_bolvanka)

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_out = Path(tmpdir) / "out.docx"
        shutil.copy2(engineer_template, tmp_out)
        doc = Document(str(tmp_out))
        _clear_body_keep_sectpr(doc)
        for line in title_lines:
            doc.add_paragraph(line)
        doc.save(str(tmp_out))
        shutil.copy2(tmp_out, output_path)

    return title_lines


def main() -> int:
    parser = argparse.ArgumentParser(description="Rebuild data/templates bolvanki from engineer template")
    parser.add_argument(
        "--engineer-template",
        type=Path,
        default=DEFAULT_ENGINEER_TEMPLATE,
        help=f"Base docx (default: {DEFAULT_ENGINEER_TEMPLATE})",
    )
    parser.add_argument(
        "--templates-dir",
        type=Path,
        default=None,
        help="Output dir (default: data/templates)",
    )
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    engineer = args.engineer_template.resolve()
    out_dir = (args.templates_dir or templates_dir()).resolve()

    if not engineer.is_file():
        print(f"[ERR] Engineer template not found: {engineer}", file=sys.stderr)
        return 1
    if not out_dir.is_dir():
        print(f"[ERR] Templates dir not found: {out_dir}", file=sys.stderr)
        return 1

    files = sorted(out_dir.glob("*.docx"))
    if not files:
        print(f"[ERR] No .docx in {out_dir}", file=sys.stderr)
        return 1

    for path in files:
        title_preview = _title_lines_from_bolvanka(path)[0][:70]
        if args.dry_run:
            print(f"[DRY] {path.name}: {title_preview!r}")
            continue
        rebuild_bolvanka(engineer, path, path)
        print(f"[OK] {path.name}: {title_preview!r}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
